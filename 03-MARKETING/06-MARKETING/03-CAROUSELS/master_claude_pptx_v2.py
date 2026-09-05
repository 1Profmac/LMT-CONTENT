from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

DARK_BG = RGBColor(0x1B, 0x2A, 0x4A)
GOLD = RGBColor(0xC8, 0xA4, 0x2C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GOLD = RGBColor(0xE8, 0xD4, 0x8C)
RED_ACCENT = RGBColor(0xCC, 0x33, 0x33)

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf

def add_para(tf, text, font_size=18, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT, space_before=Pt(6)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    p.space_before = space_before
    return p

def gold_bar(slide, top):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top), Inches(11.7), Pt(4))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

def step_badge(slide, step_num, top):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top), Inches(1.8), Inches(0.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = f"Step {step_num}"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = DARK_BG
    p.alignment = PP_ALIGN.CENTER

def series_tag(slide):
    add_text_box(slide, 0.8, 0.2, 6, 0.4, "HOW TO KEEP YOUR JOB OVER 50  |  SERIES", font_size=14, bold=True, color=GOLD)

# ─── SLIDE 1: Title ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
series_tag(slide)
gold_bar(slide, 0.55)
add_text_box(slide, 1, 1.2, 11.3, 1.5, "How to Master Claude", font_size=54, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 2.2, 11.3, 0.8, "in One Week", font_size=48, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
gold_bar(slide, 3.1)
add_text_box(slide, 1, 3.5, 11.3, 1, "A 10-Step Guide for Experienced Professionals", font_size=28, color=WHITE, alignment=PP_ALIGN.CENTER)
tf = add_text_box(slide, 1, 4.3, 11.3, 1.5, "The professionals who learn AI aren't replacing anyone.", font_size=24, color=WHITE, alignment=PP_ALIGN.CENTER)
add_para(tf, "They're becoming irreplaceable.", font_size=26, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.0, 11.3, 0.6, "Brian McKinney  |  50+TechBridge  |  #agentic50", font_size=20, color=GOLD, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.5, 11.3, 0.5, "learnmoretechnologies.com", font_size=18, color=LIGHT_GOLD, alignment=PP_ALIGN.CENTER)

# ─── SLIDE 2: The Stakes ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
series_tag(slide)
add_text_box(slide, 1, 1.0, 11.3, 1, "The Stakes Are Real", font_size=44, bold=True, color=RED_ACCENT, alignment=PP_ALIGN.CENTER)
gold_bar(slide, 1.8)
tf = add_text_box(slide, 1.5, 2.3, 10.3, 5,
    "41% of employers plan to reduce staff due to AI automation.", font_size=26, color=WHITE)
add_para(tf, "— World Economic Forum, 2025", font_size=20, color=LIGHT_GOLD)
add_para(tf, "", font_size=14)
add_para(tf, "But here's what they don't tell you:", font_size=26, color=WHITE)
add_para(tf, "", font_size=10)
add_para(tf, "The people getting cut aren't the oldest.", font_size=26, bold=True, color=WHITE)
add_para(tf, "They're the ones who never learned the tools.", font_size=26, bold=True, color=GOLD)
add_para(tf, "", font_size=14)
add_para(tf, "This isn't about age. It's about adaptability.", font_size=28, color=WHITE)
add_para(tf, "And adaptability is a skill you can learn in one week.", font_size=26, color=GOLD)

# ─── SLIDE 3: The Problem ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
series_tag(slide)
add_text_box(slide, 1, 1.0, 11.3, 1, "Most People Use Claude Wrong", font_size=44, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
gold_bar(slide, 1.8)
tf = add_text_box(slide, 1.5, 2.5, 10.3, 4.5,
    "They type a question. Get a generic answer. And wonder what the fuss is about.", font_size=26, color=WHITE)
add_para(tf, "", font_size=14)
add_para(tf, "Claude isn't a search engine. It's not a chatbot.", font_size=26, color=WHITE)
add_para(tf, "It's a digital coworker.", font_size=30, bold=True, color=GOLD)
add_para(tf, "", font_size=14)
add_para(tf, "But only if you set it up like one.", font_size=26, color=WHITE)
add_para(tf, "The difference between 'playing with AI' and getting real work done comes down to setup.", font_size=24, color=WHITE)

# ─── STEPS 1-10 ───
steps = [
    ("Download the Desktop App", [
        "Go to claude.com/download — not the website",
        "Cowork only runs inside the desktop app",
        "Mac or Windows only. No mobile. No web version.",
        "Pro account: $20/mo ($17/mo if you pay annually)",
        "",
        "That's less than most streaming subscriptions —",
        "and this one actually makes you money.",
        "",
        "CLIENT ACTION: Download and sign up before our next session."
    ]),
    ("Know Which Mode to Use", [
        "Chat  =  Quick questions. Nothing serious.",
        "Projects  =  Team use, recurring deliverables.",
        "Cowork  =  Deep solo work with actual files. YOUR WORKHORSE.",
        "Code  =  Developers only. Skip it.",
        "",
        "If you've only used Chat, you've been using",
        "the ChatGPT version of Claude. Forget it for real work.",
        "",
        "CLIENT ACTION: Open Cowork mode and explore the interface."
    ]),
    ("Build Your 4 Folders", [
        'Create one parent folder: "Claude-Cowork"',
        "",
        "ABOUT ME  —  Your identity, your writing rules",
        "PROJECTS  —  One subfolder per live project (brief + drafts + references)",
        "TEMPLATES  —  Your best past work as reusable structure",
        "CLAUDE OUTPUTS  —  The ONLY folder Claude writes to",
        "",
        "Everything else: read-only.",
        "CLIENT ACTION: Use Worksheet 1 to build your folder structure."
    ]),
    ("Write Your 2 Core Files", [
        "These two files replace 50+ prompts:",
        "",
        "about-me.md  —  Who you are, what you do,",
        "   your current priorities, what matters right now",
        "",
        'anti-ai-style.md  —  Every phrase Claude must NEVER use.',
        '   "Boundaries." "Delve." "It\'s worth noting."',
        "",
        "One great .md file beats 50 random uploads.",
        "CLIENT ACTION: Use Worksheets 2 & 3 to draft both files."
    ]),
    ("Stop Writing Prompts", [
        "Build a prompt template instead.",
        "",
        "It tells Claude to:",
        "   1. Read your ABOUT ME file first",
        "   2. Write only to CLAUDE OUTPUTS",
        "   3. Use naming convention: project_client_v1.ext",
        "",
        "Stop thinking about how to talk to Claude.",
        "Start thinking about what you need done."
    ]),
    ("Let Claude Prompt YOU", [
        "Flip the script.",
        "",
        "Instead of writing long instructions,",
        "let Claude ask the questions.",
        "",
        "Multi-select options. Drag-to-rank.",
        "You click answers in under a minute.",
        "",
        "Claude plans. You approve. It executes.",
        "You don't need to be a prompt engineer.",
        "You need to be a decision-maker. You already are."
    ]),
    ("Install One Plugin", [
        "Just one. Don't browse the whole store.",
        "",
        "Marketing  —  if you create content",
        "Data  —  if you work with spreadsheets & dashboards",
        "Legal  —  if you review contracts",
        "",
        "Shortcuts like /marketing:draft-content",
        "",
        "One plugin. Learn it. Then add another.",
        "CLIENT ACTION: Use Worksheet 4 to identify your first plugin."
    ]),
    ("Connect Your Tools", [
        "Settings > Connectors > Browse > Add",
        "",
        "Connectors let Claude act INSIDE your existing apps:",
        "   Search your Slack",
        "   Pull from your Google Docs",
        "   Reference your Notion mid-task",
        "",
        "Plugins = you use tools inside Claude",
        "Connectors = Claude reaches into YOUR tools"
    ]),
    ("Build One Project for Your Team", [
        "Shared workspace with shared context.",
        "Consistent outputs. No more copy-pasting.",
        "",
        "Even if you're solo, a Project can be your",
        "command center for recurring deliverables:",
        "",
        "   Weekly newsletter",
        "   Monthly report",
        "   Client onboarding sequence",
        "",
        "CLIENT ACTION: Use Worksheet 5 to plan your first project."
    ]),
    ("Schedule Your First Automated Task", [
        "This is the endgame.",
        "",
        "Use the /schedule plugin to set up a task",
        "that runs without you.",
        "",
        "Morning briefing. Weekly content draft.",
        "Data summary before your first cup of coffee.",
        "",
        "You wake up to a finished deliverable.",
        "That's not a tool. That's a team member",
        "who works the night shift."
    ])
]

for i, (title, bullets) in enumerate(steps, 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK_BG)
    series_tag(slide)
    step_badge(slide, i, 0.8)
    add_text_box(slide, 3, 0.75, 9.5, 0.8, title, font_size=36, bold=True, color=GOLD)
    gold_bar(slide, 1.7)
    tf = add_text_box(slide, 1.5, 2.2, 10.3, 5.0, bullets[0], font_size=24, color=WHITE)
    for bullet in bullets[1:]:
        if bullet.startswith("CLIENT ACTION"):
            add_para(tf, bullet, font_size=22, bold=True, color=LIGHT_GOLD, space_before=Pt(8))
        else:
            add_para(tf, bullet, font_size=24, color=WHITE, space_before=Pt(8))

# ─── SLIDE: Your 7-Day Plan ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
series_tag(slide)
add_text_box(slide, 1, 0.8, 11.3, 0.8, "Your 7-Day Implementation Plan", font_size=40, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
gold_bar(slide, 1.6)
tf = add_text_box(slide, 1.2, 2.0, 10.8, 5.5,
    "Day 1  —  Steps 1-2: Download app, learn the modes", font_size=24, color=WHITE)
add_para(tf, "Day 2  —  Step 3: Build your 4 folders", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Day 3  —  Step 4: Write about-me.md and anti-ai-style.md", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Day 4  —  Steps 5-6: Build prompt template, try Claude prompting you", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Day 5  —  Steps 7-8: Install one plugin, connect one tool", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Day 6  —  Step 9: Build your first project", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Day 7  —  Step 10: Schedule your first automated task", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "", font_size=14)
add_para(tf, "Use the worksheets to track your progress each day.", font_size=22, bold=True, color=GOLD)

# ─── SLIDE: Consulting CTA ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
series_tag(slide)
add_text_box(slide, 1, 1.0, 11.3, 1, "Need Help Getting Set Up?", font_size=44, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
gold_bar(slide, 1.8)
tf = add_text_box(slide, 1.5, 2.5, 10.3, 4.5,
    "50+TechBridge Consulting Options:", font_size=28, bold=True, color=WHITE)
add_para(tf, "", font_size=10)
add_para(tf, "1-on-1 Setup Session  —  We build your Claude workspace together", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Team Workshop  —  Get your whole department on Claude in one day", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "Ongoing Coaching  —  Weekly check-ins as you build AI into your workflow", font_size=24, color=WHITE, space_before=Pt(12))
add_para(tf, "", font_size=14)
add_para(tf, "Every session is customized to YOUR job, YOUR industry, YOUR goals.", font_size=24, bold=True, color=GOLD)
add_para(tf, "", font_size=14)
add_para(tf, "Book a free discovery call: learnmoretechnologies.com", font_size=26, color=GOLD)

# ─── SLIDE: Closing ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
series_tag(slide)
gold_bar(slide, 2.8)
add_text_box(slide, 1, 1.0, 11.3, 1, "The Real Point", font_size=48, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
tf = add_text_box(slide, 1, 1.8, 11.3, 1, "You're not too old. You're not too late. You're next.", font_size=30, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
tf = add_text_box(slide, 1.5, 3.3, 10.3, 1.5,
    "AI works. You just have to build the desk before you sit down at it.", font_size=28, color=WHITE, alignment=PP_ALIGN.CENTER)
add_para(tf, "", font_size=14)
add_para(tf, "One week. Ten steps. You'll never go back.", font_size=30, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
add_para(tf, "", font_size=20)
add_para(tf, "Follow #agentic50 for weekly AI guidance built for experienced professionals.", font_size=22, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.2, 11.3, 0.6, "learnmoretechnologies.com", font_size=20, color=GOLD, alignment=PP_ALIGN.CENTER)

outpath = r"C:\Users\Jordyn\Desktop\LMT-CONTENT\06-MARKETING\Master_Claude_One_Week_v2.pptx"
prs.save(outpath)
print(f"PowerPoint saved: {outpath}")
