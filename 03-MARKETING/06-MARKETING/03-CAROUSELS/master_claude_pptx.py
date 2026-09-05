from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Brand colors
DARK_BG = RGBColor(0x1B, 0x2A, 0x4A)
GOLD = RGBColor(0xC8, 0xA4, 0x2C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF0, 0xF0, 0xF0)
BLACK = RGBColor(0x10, 0x10, 0x10)

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf

def add_para(tf, text, font_size=18, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT, space_before=Pt(6)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    p.space_before = space_before
    return p

def add_gold_bar(slide, top):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top), Inches(11.7), Pt(4))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

def add_step_number(slide, step_num, top):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top), Inches(1.8), Inches(0.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = f"Step {step_num}"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = DARK_BG
    p.alignment = PP_ALIGN.CENTER

# ─── SLIDE 1: Title ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
add_gold_bar(slide, 2.8)
add_text_box(slide, 1, 1.2, 11.3, 1.5, "How to Master Claude", font_size=54, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 2.2, 11.3, 0.8, "in One Week", font_size=48, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 3.5, 11.3, 1, "A 10-Step Guide for Professionals Over 50", font_size=28, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 5.5, 11.3, 0.6, "Brian McKinney  |  50+TechBridge  |  #agentic50", font_size=20, color=GOLD, alignment=PP_ALIGN.CENTER)

# ─── SLIDE 2: The Problem ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
add_gold_bar(slide, 1.8)
add_text_box(slide, 1, 0.8, 11.3, 1, "Most People Use Claude Wrong", font_size=44, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
tf = add_text_box(slide, 1.5, 2.5, 10.3, 4.5,
    "They type a question. Get a generic answer. And wonder what the fuss is about.", font_size=26, color=WHITE)
add_para(tf, "", font_size=14)
add_para(tf, "Claude isn't a search engine. It's not a chatbot.", font_size=26, color=WHITE)
add_para(tf, "It's a digital coworker.", font_size=30, bold=True, color=GOLD)
add_para(tf, "", font_size=14)
add_para(tf, "But only if you set it up like one.", font_size=26, color=WHITE)
add_para(tf, "The difference between 'playing with AI' and getting real work done comes down to setup.", font_size=24, color=WHITE)

# ─── SLIDES 3-12: Steps 1-10 ───
steps = [
    {
        "title": "Download the Desktop App",
        "bullets": [
            "Go to claude.com/download — not the website",
            "Cowork only runs inside the desktop app",
            "Mac or Windows only. No mobile. No web version.",
            "Pro account: $20/mo ($17/mo if you pay annually)",
            "That's less than most streaming subscriptions —",
            "and this one actually makes you money."
        ]
    },
    {
        "title": "Know Which Mode to Use",
        "bullets": [
            "Chat  =  Quick questions. Nothing serious.",
            "Projects  =  Team use, recurring deliverables.",
            "Cowork  =  Deep solo work with actual files. YOUR WORKHORSE.",
            "Code  =  Developers only. Skip it.",
            "",
            "If you've only used Chat, you've been using",
            "the ChatGPT version of Claude. Forget it for real work."
        ]
    },
    {
        "title": "Build Your 4 Folders",
        "bullets": [
            'Create one parent folder: "Claude-Cowork"',
            "",
            "ABOUT ME  —  Your identity, your writing rules",
            "PROJECTS  —  One subfolder per live project (brief + drafts + references)",
            "TEMPLATES  —  Your best past work as reusable structure",
            "CLAUDE OUTPUTS  —  The ONLY folder Claude writes to",
            "",
            "Everything else: read-only."
        ]
    },
    {
        "title": "Write Your 2 Core Files",
        "bullets": [
            "These two files replace 50+ prompts:",
            "",
            "about-me.md  —  Who you are, what you do,",
            "   your current priorities, what matters right now",
            "",
            'anti-ai-style.md  —  Every phrase Claude must NEVER use.',
            '   "Boundaries." "Delve." "It\'s worth noting."',
            "",
            "One great .md file beats 50 random uploads."
        ]
    },
    {
        "title": "Stop Writing Prompts",
        "bullets": [
            "Build a prompt template instead.",
            "",
            "It tells Claude to:",
            "   1. Read your ABOUT ME file first",
            "   2. Write only to CLAUDE OUTPUTS",
            "   3. Use naming convention: project_client_v1.ext",
            "",
            "Stop thinking about how to talk to Claude.",
            "Start thinking about what you need done."
        ]
    },
    {
        "title": "Let Claude Prompt YOU",
        "bullets": [
            "Flip the script.",
            "",
            "Instead of writing long instructions,",
            "let Claude ask the questions.",
            "",
            "Multi-select options. Drag-to-rank.",
            "You click answers in under a minute.",
            "",
            "Claude plans. You approve. It executes.",
            "You don't need to be a prompt engineer.",
            "You need to be a decision-maker. You already are."
        ]
    },
    {
        "title": "Install One Plugin",
        "bullets": [
            "Just one. Don't browse the whole store.",
            "",
            "Marketing  —  if you create content",
            "Data  —  if you work with spreadsheets & dashboards",
            "Legal  —  if you review contracts",
            "",
            "Shortcuts like /marketing:draft-content",
            "",
            "One plugin. Learn it. Then add another."
        ]
    },
    {
        "title": "Connect Your Tools",
        "bullets": [
            "Settings > Connectors > Browse > Add",
            "",
            "Connectors let Claude act INSIDE your existing apps:",
            "   Search your Slack",
            "   Pull from your Google Docs",
            "   Reference your Notion mid-task",
            "",
            "Plugins = you use tools inside Claude",
            "Connectors = Claude reaches into YOUR tools"
        ]
    },
    {
        "title": "Build One Project for Your Team",
        "bullets": [
            "Shared workspace with shared context.",
            "Consistent outputs. No more copy-pasting.",
            "",
            "Even if you're solo, a Project can be your",
            "command center for recurring deliverables:",
            "",
            "   Weekly newsletter",
            "   Monthly report",
            "   Client onboarding sequence"
        ]
    },
    {
        "title": "Schedule Your First Automated Task",
        "bullets": [
            "This is the endgame.",
            "",
            "Use the /schedule plugin to set up a task",
            "that runs without you.",
            "",
            "Morning briefing. Weekly content draft.",
            "Data summary before your first cup of coffee.",
            "",
            "You wake up to a finished deliverable.",
            "That's not a tool. That's a team member",
            "who works the night shift."
        ]
    }
]

for i, step in enumerate(steps, 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK_BG)
    add_step_number(slide, i, 0.6)
    add_text_box(slide, 3, 0.55, 9.5, 0.8, step["title"], font_size=36, bold=True, color=GOLD)
    add_gold_bar(slide, 1.5)

    tf = add_text_box(slide, 1.5, 2.0, 10.3, 5.0, step["bullets"][0], font_size=24, color=WHITE)
    for bullet in step["bullets"][1:]:
        add_para(tf, bullet, font_size=24, color=WHITE, space_before=Pt(8))

# ─── SLIDE 13: Closing ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)
add_gold_bar(slide, 2.8)
add_text_box(slide, 1, 1.0, 11.3, 1, "The Real Point", font_size=48, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
tf = add_text_box(slide, 1.5, 3.2, 10.3, 1.5,
    "AI works. You just have to build the desk before you sit down at it.", font_size=28, color=WHITE, alignment=PP_ALIGN.CENTER)
add_para(tf, "", font_size=14)
add_para(tf, "One week. Ten steps. You'll never go back.", font_size=30, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
add_para(tf, "", font_size=20)
add_para(tf, "Follow #agentic50 for weekly AI guidance built for experienced professionals.", font_size=22, color=WHITE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.2, 11.3, 0.6, "learnmoretechnologies.com", font_size=20, color=GOLD, alignment=PP_ALIGN.CENTER)

outpath = r"C:\Users\Jordyn\Desktop\LMT-CONTENT\06-MARKETING\Master_Claude_One_Week.pptx"
prs.save(outpath)
print(f"PowerPoint saved: {outpath}")
