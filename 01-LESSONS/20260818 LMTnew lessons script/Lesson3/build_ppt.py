"""
Build Lesson 3 PowerPoint — Don't Get Scammed
Brand: Navy #0E1C2F | Gold #C8942E | White text | DM Sans (falls back to Calibri)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

# ── Colors ──────────────────────────────────────────────────────────────────
NAVY  = RGBColor(0x0E, 0x1C, 0x2F)
GOLD  = RGBColor(0xC8, 0x94, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xCC, 0xD6, 0xE0)   # muted blue-white for body text

FONT_FACE = "DM Sans"

# ── Slide dimensions (16:9 widescreen) ──────────────────────────────────────
W = Inches(13.333)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]   # truly blank layout

# ── Helper: add a shape with solid fill ─────────────────────────────────────
def rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

# ── Helper: add a text box ───────────────────────────────────────────────────
def textbox(slide, left, top, width, height, text,
            font_size=28, bold=False, color=WHITE,
            align=PP_ALIGN.LEFT, font=FONT_FACE, wrap=True):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txb

# ── Helper: multi-line bullet block ─────────────────────────────────────────
def bullets(slide, left, top, width, height, items,
            font_size=22, color=WHITE, bullet_char="→", font=FONT_FACE):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf  = txb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = f"{bullet_char}  {item}"
        run.font.name = font
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
    return txb

# ── Helper: gold rule line ───────────────────────────────────────────────────
def gold_rule(slide, top, left=Inches(0.6), width=Inches(12.13)):
    rect(slide, left, top, width, Inches(0.04), GOLD)

# ── Helper: navy background ──────────────────────────────────────────────────
def bg(slide):
    rect(slide, 0, 0, W, H, NAVY)

# ── Helper: slide number tag ─────────────────────────────────────────────────
def tag(slide, label, num=None):
    text = f"{label}  {num}" if num else label
    textbox(slide, Inches(11.8), Inches(7.1), Inches(1.4), Inches(0.35),
            text, font_size=10, color=GOLD, align=PP_ALIGN.RIGHT)

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, Inches(5.6), W, Inches(1.9), RGBColor(0x07, 0x10, 0x1A))  # darker footer

# Gold side bar
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.6), Inches(10), Inches(0.6),
        "LESSON 3", font_size=16, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(1.2), Inches(11.5), Inches(2.2),
        "Don't Get Scammed", font_size=60, bold=True, color=WHITE)

textbox(s, Inches(0.5), Inches(3.4), Inches(11), Inches(0.7),
        "3 Skills That Keep Scammers Away From Your Money",
        font_size=24, color=LIGHT)

gold_rule(s, Inches(4.3))

textbox(s, Inches(0.5), Inches(4.5), Inches(6), Inches(0.5),
        "50+TechBridge  ·  Brian McKinney", font_size=16, color=GOLD)

textbox(s, Inches(0.5), Inches(5.0), Inches(9), Inches(0.5),
        "learnmoretechnologies.com/start-free-lessons", font_size=14, color=LIGHT)

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Hook / Setup
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.5), Inches(12), Inches(1.2),
        "You already did the brave part.", font_size=36, bold=True, color=WHITE)

textbox(s, Inches(0.5), Inches(1.7), Inches(11.5), Inches(0.8),
        "You asked a machine a real question. You stayed in charge of the answer.",
        font_size=22, color=LIGHT)

gold_rule(s, Inches(2.7))

textbox(s, Inches(0.5), Inches(2.9), Inches(11), Inches(1.0),
        "This lesson is how you stay in charge when a stranger wants your fear.",
        font_size=24, color=WHITE)

# Big quote
rect(s, Inches(0.5), Inches(4.1), Inches(12.3), Inches(1.6), RGBColor(0x16, 0x2A, 0x44))
textbox(s, Inches(0.75), Inches(4.2), Inches(11.5), Inches(1.4),
        '"Frightened people send money.  Clear people hang up."',
        font_size=28, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

textbox(s, Inches(0.5), Inches(6.1), Inches(5), Inches(0.5),
        "Three skills.  That is the whole class.", font_size=18, color=LIGHT)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 3 — The Real Numbers
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(6), Inches(0.5),
        "THE REAL NUMBERS", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11.5), Inches(0.9),
        "You are not the punchline.", font_size=34, bold=True, color=WHITE)

gold_rule(s, Inches(1.85))

# Two stat boxes
for left, text, sub in [
    (Inches(0.5),  "$2.4 Billion",  "reported to FTC by adults 60+  (2024)"),
    (Inches(6.8),  "~$5 Billion",   "reported to FBI Internet Crime Center"),
]:
    rect(s, left, Inches(2.0), Inches(6.0), Inches(1.8), RGBColor(0x16, 0x2A, 0x44))
    textbox(s, left + Inches(0.2), Inches(2.1), Inches(5.6), Inches(0.85),
            text, font_size=36, bold=True, color=GOLD)
    textbox(s, left + Inches(0.2), Inches(2.9), Inches(5.6), Inches(0.7),
            sub, font_size=14, color=LIGHT)

# The hero stat
rect(s, Inches(0.5), Inches(4.05), Inches(12.3), Inches(1.7), RGBColor(0x07, 0x10, 0x1A))
textbox(s, Inches(0.75), Inches(4.1), Inches(11.5), Inches(0.65),
        "The line almost nobody reads out loud:", font_size=17, color=LIGHT, align=PP_ALIGN.CENTER)
textbox(s, Inches(0.75), Inches(4.65), Inches(11.5), Inches(0.9),
        "3 in 4 older adults who reported a scam lost NO money.  They spotted it.",
        font_size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Skill 1 Title Card
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, W, Inches(0.18), GOLD)          # top gold bar

textbox(s, Inches(0.5), Inches(1.2), Inches(4), Inches(0.8),
        "SKILL 1", font_size=20, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(1.9), Inches(11), Inches(1.6),
        "Spot a Fake", font_size=72, bold=True, color=WHITE)

gold_rule(s, Inches(3.7))

textbox(s, Inches(0.5), Inches(3.9), Inches(11), Inches(1.0),
        "Look at the costume, not the logo.", font_size=30, color=LIGHT)

textbox(s, Inches(0.5), Inches(5.2), Inches(11), Inches(0.7),
        "A fake wants you to believe the costume.", font_size=22, color=LIGHT)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Spot a Fake: The 4 Tells
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(6), Inches(0.5),
        "SKILL 1  ·  SPOT A FAKE", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11), Inches(0.7),
        "Four things to check — every time.", font_size=28, bold=True, color=WHITE)

gold_rule(s, Inches(1.7))

tells = [
    ("THE SENDER",   "Would Amazon write from  email-secure-alert.com?  Hover and look."),
    ("THE GREETING", '"Dear Customer" from a bank that knows your name is a tell.'),
    ("THE ASK",      "A link. A phone number. A QR code. A download. A code they want you to read."),
    ("THE SURPRISE", "You did not start this conversation. That matters more than the grammar.\nAI fixed the spelling. 'No typos' is no longer your safety test."),
]

top = Inches(1.9)
for label, body in tells:
    rect(s, Inches(0.5), top, Inches(0.08), Inches(0.55), GOLD)
    textbox(s, Inches(0.72), top - Inches(0.02), Inches(2.8), Inches(0.45),
            label, font_size=14, bold=True, color=GOLD)
    textbox(s, Inches(0.72), top + Inches(0.35), Inches(12.0), Inches(0.65),
            body, font_size=17, color=WHITE)
    top += Inches(1.28)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Common Costumes
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(6), Inches(0.5),
        "SKILL 1  ·  COMMON COSTUMES", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11), Inches(0.7),
        "Treat these as a costume until YOU started the call.", font_size=26, bold=True, color=WHITE)

gold_rule(s, Inches(1.7))

costumes = ["Your Bank", "Medicare", "Amazon", "Microsoft / Apple",
            "IRS or Social Security", "A Sheriff or Police Department",
            "A Grandchild in Trouble", "Tech Support That Called You"]

col_w = Inches(5.8)
for i, name in enumerate(costumes):
    col = i % 2
    row = i // 2
    left = Inches(0.5) + col * col_w
    top  = Inches(2.0) + row * Inches(1.1)
    rect(s, left, top, col_w - Inches(0.2), Inches(0.85), RGBColor(0x16, 0x2A, 0x44))
    textbox(s, left + Inches(0.15), top + Inches(0.15), col_w - Inches(0.4), Inches(0.6),
            name, font_size=20, bold=True, color=WHITE)

# Rule at bottom
rect(s, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.75), RGBColor(0x07, 0x10, 0x1A))
textbox(s, Inches(0.7), Inches(6.4), Inches(12), Inches(0.6),
        "Hang up.  Look up the real number.  Call that number — not the one in the text.",
        font_size=18, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Skill 2 Title Card
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, W, Inches(0.18), GOLD)

textbox(s, Inches(0.5), Inches(1.2), Inches(4), Inches(0.8),
        "SKILL 2", font_size=20, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(1.9), Inches(11), Inches(1.6),
        "Name the Pressure", font_size=66, bold=True, color=WHITE)

gold_rule(s, Inches(3.7))

textbox(s, Inches(0.5), Inches(3.9), Inches(11), Inches(1.0),
        "Fakes work because of feelings — not because you cannot read.", font_size=26, color=LIGHT)

textbox(s, Inches(0.5), Inches(5.1), Inches(11), Inches(0.7),
        "Learn the four levers.  Say them out loud.", font_size=22, color=GOLD)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 8 — The 4 Levers
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(6), Inches(0.5),
        "SKILL 2  ·  THE 4 LEVERS", font_size=14, bold=True, color=GOLD)

levers = [
    ("HURRY",     "20 minutes. A warrant. The refund expires.\nIf you are not allowed to think, you are not in a real official process."),
    ("FEAR",      "Arrest. Deportation. A virus. A stolen identity they will help you fix.\nFear makes smart people obedient."),
    ("SECRECY",   "\"Don't tell your wife. She would worry. Don't tell the bank teller.\"\nA real emergency can stand a second adult."),
    ("AUTHORITY", "A badge number. A case number. A familiar logo.\nCaller ID can be faked. The badge can be faked."),
]

col_w = Inches(6.1)
for i, (label, body) in enumerate(levers):
    col = i % 2
    row = i // 2
    left = Inches(0.5) + col * col_w
    top  = Inches(1.05) + row * Inches(2.85)
    rect(s, left, top, col_w - Inches(0.2), Inches(2.6), RGBColor(0x16, 0x2A, 0x44))
    textbox(s, left + Inches(0.2), top + Inches(0.15), col_w - Inches(0.5), Inches(0.55),
            label, font_size=22, bold=True, color=GOLD)
    textbox(s, left + Inches(0.2), top + Inches(0.7), col_w - Inches(0.5), Inches(1.7),
            body, font_size=17, color=WHITE)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 9 — Never Move Money
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(8), Inches(0.5),
        "SKILL 2  ·  THE FIFTH LEVER", font_size=14, bold=True, color=GOLD)

# Big warning
rect(s, Inches(0.5), Inches(0.9), Inches(12.3), Inches(1.4), RGBColor(0xC8, 0x94, 0x2E))
textbox(s, Inches(0.7), Inches(0.95), Inches(12), Inches(1.2),
        "NEVER move money to protect it.", font_size=38, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

textbox(s, Inches(0.5), Inches(2.5), Inches(12), Inches(0.9),
        "Danger payment methods — if you hear any of these, stop and hang up:",
        font_size=20, color=LIGHT)

danger = ["Gift cards (Walmart, Amazon, Google Play)",
          "Wire transfers",
          "Crypto / Bitcoin ATM",
          "A courier at your door collecting cash",
          "Remote access to your computer"]

bullets(s, Inches(0.8), Inches(3.4), Inches(11.5), Inches(3.0),
        danger, font_size=20, color=WHITE, bullet_char="✕")

rect(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.7), RGBColor(0x07, 0x10, 0x1A))
textbox(s, Inches(0.7), Inches(6.55), Inches(12), Inches(0.55),
        "No legitimate tax office, bank, or grandchild-in-trouble story needs a Walmart gift card.",
        font_size=17, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 10 — AI Voice Cloning / Family Code Word
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(8), Inches(0.5),
        "SKILL 2  ·  THE NEW COSTUME", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11.5), Inches(0.9),
        "A voice you love — cloned by AI.", font_size=32, bold=True, color=WHITE)

gold_rule(s, Inches(1.85))

textbox(s, Inches(0.5), Inches(2.0), Inches(11.5), Inches(0.7),
        "AI-enabled scams (realistic voices and video) jumped ~20× from 2023 to 2025.\nA familiar voice is no longer proof.",
        font_size=20, color=LIGHT)

# Code word box
rect(s, Inches(0.5), Inches(3.1), Inches(12.3), Inches(2.6), RGBColor(0x16, 0x2A, 0x44))
textbox(s, Inches(0.7), Inches(3.2), Inches(11.5), Inches(0.6),
        "A 20th-century fix for a 21st-century trick:", font_size=18, color=GOLD)
textbox(s, Inches(0.7), Inches(3.75), Inches(11.5), Inches(0.7),
        "Pick a family code word.", font_size=30, bold=True, color=WHITE)
textbox(s, Inches(0.7), Inches(4.4), Inches(11.5), Inches(1.1),
        "Known only to your people. Used when money or panic shows up.\nA word you would never post on Facebook. Share it this week.",
        font_size=18, color=LIGHT)

textbox(s, Inches(0.5), Inches(5.9), Inches(12.3), Inches(0.7),
        "If a panicked voice asks for bail, tickets, or a gift card and cannot give the word — hang up.",
        font_size=19, bold=True, color=WHITE)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 11 — Skill 3 Title Card
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, W, Inches(0.18), GOLD)

textbox(s, Inches(0.5), Inches(1.2), Inches(4), Inches(0.8),
        "SKILL 3", font_size=20, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(1.9), Inches(12), Inches(2.0),
        "Passwords &\nthe Second Lock", font_size=56, bold=True, color=WHITE)

gold_rule(s, Inches(4.1))

textbox(s, Inches(0.5), Inches(4.3), Inches(11), Inches(0.7),
        "A password is a key.  Most people have one key and twenty doors.", font_size=24, color=LIGHT)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 12 — Better Passwords
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(8), Inches(0.5),
        "SKILL 3  ·  BETTER PASSWORDS", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11.5), Inches(0.7),
        "NIST (the people who write the grown-up rules) changed the advice.", font_size=24, bold=True, color=WHITE)

gold_rule(s, Inches(1.7))

rules = [
    "LONGER is better — aim for 15 or more characters",
    "A PHRASE beats a scrambled string:  maple-porch-radio-Thursday",
    "You do NOT have to change it every 90 days",
    "UNIQUE matters — email password cannot match bank password",
    "Email is the skeleton key — if they get email, they reset everything else",
]

bullets(s, Inches(0.6), Inches(2.0), Inches(12.3), Inches(3.8),
        rules, font_size=20, color=WHITE)

# Two honest options
rect(s, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.7), RGBColor(0x16, 0x2A, 0x44))
textbox(s, Inches(0.7), Inches(5.6), Inches(12), Inches(0.5),
        "Can't remember 20 phrases?  Two honest options:", font_size=17, bold=True, color=GOLD)
textbox(s, Inches(0.7), Inches(6.05), Inches(12), Inches(0.9),
        "1.  A password manager (vault app) — get help setting it up once.\n"
        "2.  A small notebook in a drawer or lockbox at home. Not shameful. Better than one reused password.",
        font_size=16, color=WHITE)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 13 — Two-Factor Authentication
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(8), Inches(0.5),
        "SKILL 3  ·  THE SECOND LOCK", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11), Inches(0.75),
        "Two-Factor Authentication (2FA)", font_size=34, bold=True, color=WHITE)

textbox(s, Inches(0.5), Inches(1.65), Inches(11), Inches(0.55),
        "Password  +  something you have.  A thief with only the password should bounce.",
        font_size=20, color=LIGHT)

gold_rule(s, Inches(2.3))

steps = [
    "Turn it on for EMAIL first — then your bank or credit union",
    "Where to find it:  Settings → Security → Two-factor / 2-step / Multifactor",
    "Any second factor is better than none  (CISA's own words)",
    "An authenticator app or passkey is stronger than a text code",
    "Text codes can be stolen — but still: turn it on",
]

bullets(s, Inches(0.6), Inches(2.5), Inches(12.3), Inches(3.5),
        steps, font_size=20, color=WHITE)

# Warning
rect(s, Inches(0.5), Inches(6.1), Inches(12.3), Inches(1.1), RGBColor(0xC8, 0x94, 0x2E))
textbox(s, Inches(0.7), Inches(6.15), Inches(12), Inches(0.9),
        "NEVER read a code to anyone — not even someone who says they are the bank.",
        font_size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 14 — If It Already Happened
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(6), Inches(0.5),
        "IF IT ALREADY HAPPENED", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11.5), Inches(0.7),
        "This is not a moral failure.", font_size=32, bold=True, color=WHITE)

textbox(s, Inches(0.5), Inches(1.6), Inches(11.5), Inches(0.6),
        "Professionals ran a play on a human nervous system. You are still a competent person.",
        font_size=20, color=LIGHT)

gold_rule(s, Inches(2.3))

steps = [
    "STOP sending money immediately",
    "Call the bank — number on the back of your card — say wire / gift card / crypto",
    "Change your email password and turn on 2FA right now",
    "Report at  reportfraud.ftc.gov  and  ic3.gov",
    "Tell one trusted person — speed matters, silence guarantees no recovery",
    "If you gave a code or password — change it from a device you trust",
    "Identity theft steps:  identitytheft.gov",
]

bullets(s, Inches(0.6), Inches(2.5), Inches(12.3), Inches(4.5),
        steps, font_size=19, color=WHITE)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 15 — Homework (5 Steps)
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)

textbox(s, Inches(0.5), Inches(0.4), Inches(6), Inches(0.5),
        "YOUR HOMEWORK  ·  20 MINUTES", font_size=14, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(0.85), Inches(11), Inches(0.7),
        "Five things.  Do them this week.", font_size=30, bold=True, color=WHITE)

gold_rule(s, Inches(1.65))

hw = [
    ("1", "Put the hang-up rule on the fridge:",
          "Unexpected + hurry + money or a code  =  hang up and verify"),
    ("2", "Send the family code word to the people who matter.  Not posted online.", ""),
    ("3", "Turn on email two-factor authentication.  Then the bank.", ""),
    ("4", "Set one unique long passphrase for email if it matches other sites.", ""),
    ("5", "If a scam already happened — file both reports today.", ""),
]

top = Inches(1.85)
for num, main, sub in hw:
    rect(s, Inches(0.5), top, Inches(0.55), Inches(0.55), GOLD)
    textbox(s, Inches(0.5), top + Inches(0.04), Inches(0.55), Inches(0.5),
            num, font_size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1.2), top, Inches(11.5), Inches(0.5),
            main, font_size=19, bold=True, color=WHITE)
    if sub:
        textbox(s, Inches(1.2), top + Inches(0.44), Inches(11.5), Inches(0.4),
                sub, font_size=15, color=GOLD)
        top += Inches(0.95)
    else:
        top += Inches(0.72)

tag(s, "50+TechBridge · Lesson 3")

# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 16 — Closing
# ═══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, Inches(0.18), H, GOLD)
rect(s, Inches(0.18), Inches(5.3), W, Inches(2.2), RGBColor(0x07, 0x10, 0x1A))

textbox(s, Inches(0.5), Inches(0.6), Inches(11), Inches(0.6),
        "You now have the three-lesson arc we promised.", font_size=22, color=LIGHT)

textbox(s, Inches(0.5), Inches(1.2), Inches(11.5), Inches(3.0),
        "You can spot a costume.\nYou can name the pressure.\nYou can lock the door.",
        font_size=36, bold=True, color=WHITE)

gold_rule(s, Inches(4.2))

# Tagline
textbox(s, Inches(0.5), Inches(4.35), Inches(11.5), Inches(0.75),
        "Hang up.  Look up.  Call back.",
        font_size=34, bold=True, color=GOLD)

textbox(s, Inches(0.5), Inches(5.45), Inches(5), Inches(0.5),
        "Brian McKinney  ·  50+TechBridge", font_size=16, color=LIGHT)

textbox(s, Inches(0.5), Inches(5.9), Inches(9), Inches(0.4),
        "Free lessons:  learnmoretechnologies.com/start-free-lessons", font_size=14, color=LIGHT)

textbox(s, Inches(0.5), Inches(6.3), Inches(9), Inches(0.4),
        "Report a scam:  reportfraud.ftc.gov  ·  ic3.gov", font_size=14, color=LIGHT)

tag(s, "50+TechBridge · Lesson 3")

# ── Save ─────────────────────────────────────────────────────────────────────
out = r"C:\Users\Jordyn\Desktop\LMT-CONTENT\01-LESSONS\20260818 LMTnew lessons script\Lesson3\Dont-Get-Scammed-Lesson3.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
