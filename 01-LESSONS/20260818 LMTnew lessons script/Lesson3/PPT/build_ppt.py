"""
Lesson 3 — Don't Get Scammed
50+TechBridge classroom / HeyGen deck
16:9  ·  Navy #0E1C2F  ·  Gold #C8942E
Large type for 50+ readers.

Run:  python build_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Brand (match Lesson 2) ──────────────────────────────────────────────────
NAVY      = RGBColor(0x0E, 0x1C, 0x2F)
NAVY_CARD = RGBColor(0x16, 0x2C, 0x45)
GOLD      = RGBColor(0xC8, 0x94, 0x2E)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT     = RGBColor(0xC8, 0xCD, 0xD6)
MUTED     = RGBColor(0x9A, 0xA3, 0xB0)
RED       = RGBColor(0xE8, 0x5D, 0x4C)
GREEN     = RGBColor(0x4C, 0xB8, 0x8A)
FONT      = "Calibri"

W = Inches(13.333)
H = Inches(7.5)

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "DontGetScammed-Lesson3.pptx")


def new_prs():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY
    return slide


def _run(p, text, size, bold=False, color=WHITE, italic=False):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    return run


def add_textbox(slide, text, left, top, width, height,
                font_size=28, bold=False, color=WHITE,
                align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {
            MSO_ANCHOR.TOP: "t",
            MSO_ANCHOR.MIDDLE: "ctr",
            MSO_ANCHOR.BOTTOM: "b",
        }.get(anchor, "t"))
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, font_size, bold, color, italic)
    return box


def add_rect(slide, left, top, width, height, fill_color, line_color=None, line_pt=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_pt)
    return shape


def add_round(slide, left, top, width, height, fill_color, line_color=None, line_pt=2):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_pt)
    # Softer corners
    try:
        adj = shape.adjustments
        adj[0] = 0.08
    except Exception:
        pass
    return shape


def chrome(slide, footer="50+TechBridge  ·  Brian McKinney  ·  Lesson 3 of 3"):
    """Gold bars + series tag + footer — Lesson 2 language."""
    add_textbox(slide, "DIGITAL PIONEER  |  LESSON 3",
                Inches(0.55), Inches(0.18), Inches(12.2), Inches(0.35),
                font_size=14, bold=True, color=GOLD, align=PP_ALIGN.LEFT)
    add_rect(slide, Inches(0.55), Inches(0.52), Inches(12.23), Pt(3.5), GOLD)
    add_rect(slide, Inches(0.55), Inches(7.05), Inches(12.23), Pt(3.5), GOLD)
    add_textbox(slide, footer,
                Inches(0.55), Inches(7.12), Inches(12.23), Inches(0.32),
                font_size=13, bold=False, color=GOLD, align=PP_ALIGN.CENTER)


def title_bar(slide, title, color=GOLD):
    add_textbox(slide, title,
                Inches(0.55), Inches(0.62), Inches(12.23), Inches(0.7),
                font_size=32, bold=True, color=color)


def bullets(slide, items, left, top, width, height, space_before=10):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, str):
            text, size, bold, color, italic = item, 24, False, WHITE, False
        else:
            text = item[0]
            size = item[1] if len(item) > 1 else 24
            bold = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 else WHITE
            italic = item[4] if len(item) > 4 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(space_before)
        _run(p, text, size, bold, color, italic)
    return box


def card_text(shape, lines, font_size=20, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.16)
    tf.margin_bottom = Inches(0.12)
    first = True
    for line in lines:
        if isinstance(line, str):
            text, size, b, c = line, font_size, bold, color
        else:
            text = line[0]
            size = line[1] if len(line) > 1 else font_size
            b = line[2] if len(line) > 2 else bold
            c = line[3] if len(line) > 3 else color
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(4)
        _run(p, text, size, b, c)
    return tf


# ============================================================================
#  SLIDES
# ============================================================================

prs = new_prs()

# ── 1. Title ────────────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
add_textbox(s, "LESSON 3 OF 3",
            Inches(0.55), Inches(1.7), Inches(12.23), Inches(0.45),
            font_size=20, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Don't Get Scammed",
            Inches(0.55), Inches(2.2), Inches(12.23), Inches(1.3),
            font_size=60, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "Three skills that keep you in charge",
            Inches(0.55), Inches(3.55), Inches(12.23), Inches(0.6),
            font_size=28, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Frightened people send money.  Clear people hang up.",
            Inches(0.55), Inches(4.5), Inches(12.23), Inches(0.55),
            font_size=22, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)
add_textbox(s, "Brian McKinney  ·  50+TechBridge",
            Inches(0.55), Inches(5.9), Inches(12.23), Inches(0.45),
            font_size=20, color=MUTED, align=PP_ALIGN.CENTER)

# ── 2. Opening ──────────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "You already did the brave part.")
bullets(s, [
    ("You asked a machine a real question", 28, False, WHITE),
    ("and stayed in charge of the answer.", 28, False, WHITE),
    ("", 10, False, WHITE),
    ("This lesson is how you stay in charge", 26, False, LIGHT),
    ("when a stranger wants your fear.", 26, False, LIGHT),
    ("", 12, False, WHITE),
    ("We will not treat you like a child.", 26, True, GOLD),
    ("We will treat you like someone who already has judgment —", 24, False, WHITE),
    ("and who is being hunted by professionals.", 24, False, WHITE),
], Inches(0.7), Inches(1.45), Inches(12.0), Inches(5.3), space_before=6)

# ── 3. Three skills ─────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "Three skills. That is the whole class.")
cards = [
    ("1", "Spot a Fake", "Look at the costume,\nnot the logo."),
    ("2", "Name the Pressure", "Hurry. Fear.\nSecrecy. Authority."),
    ("3", "Lock the Door", "A long passphrase\nand a second lock."),
]
for i, (num, head, body) in enumerate(cards):
    left = Inches(0.55 + i * 4.2)
    card = add_round(s, left, Inches(1.7), Inches(3.95), Inches(4.5),
                     NAVY_CARD, GOLD, 2.5)
    add_textbox(s, num,
                left, Inches(1.95), Inches(3.95), Inches(0.7),
                font_size=36, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(s, head,
                left + Inches(0.15), Inches(2.7), Inches(3.65), Inches(1.0),
                font_size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(s, body,
                left + Inches(0.2), Inches(3.85), Inches(3.55), Inches(1.6),
                font_size=20, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 4. You are not the punchline ────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "You are not the punchline.")
# Two stat cards
c1 = add_round(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(2.55),
               NAVY_CARD, GOLD, 2)
card_text(c1, [
    ("FTC  ·  adults 60+", 16, True, GOLD),
    ("$2.4 billion", 40, True, WHITE),
    ("reported lost to fraud in 2024", 20, False, LIGHT),
    ("about 4× what they reported in 2020", 16, False, MUTED),
], align=PP_ALIGN.CENTER)
c2 = add_round(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(2.55),
               NAVY_CARD, GOLD, 2)
card_text(c2, [
    ("FBI  ·  internet crime center", 16, True, GOLD),
    ("Nearly $5 billion", 36, True, WHITE),
    ("lost by people over 60", 20, False, LIGHT),
    ("most complaints of any age group", 16, False, MUTED),
], align=PP_ALIGN.CENTER)
add_textbox(s, "Those numbers are real. They are ugly.",
            Inches(0.55), Inches(4.4), Inches(12.23), Inches(0.5),
            font_size=24, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "They are not a verdict on your intelligence.",
            Inches(0.55), Inches(4.95), Inches(12.23), Inches(0.55),
            font_size=28, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Do not mix FTC and FBI totals. Different stacks of reports.",
            Inches(0.55), Inches(5.7), Inches(12.23), Inches(0.4),
            font_size=16, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

# ── 5. Three in four ────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "The line almost nobody reads out loud")
add_textbox(s, "About 3 in 4",
            Inches(0.55), Inches(1.6), Inches(12.23), Inches(1.1),
            font_size=64, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
add_textbox(s, "older adults who reported a scam to the FTC lost no money.",
            Inches(0.7), Inches(2.75), Inches(11.9), Inches(0.7),
            font_size=26, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "They spotted it. They saw it. They did not pay.",
            Inches(0.7), Inches(3.5), Inches(11.9), Inches(0.55),
            font_size=26, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "The dollars exploded because a smaller number of scams now take\n"
               "life-changing amounts — $100,000 or more — through fake investments,\n"
               "fake banks, fake government, and fake tech support.",
            Inches(0.8), Inches(4.4), Inches(11.7), Inches(1.6),
            font_size=20, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 6. Phone vs social ──────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "Phone still takes the most money per person")
c1 = add_round(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(3.4),
               NAVY_CARD, RED, 2.5)
card_text(c1, [
    ("THE CALL", 16, True, RED),
    ("About $2,200", 36, True, WHITE),
    ("median loss on a phone-started\nfraud for older adults (FTC)", 18, False, LIGHT),
    ("", 10, False, WHITE),
    ("If you are 80 and the phone rings,\nthat call is still the dangerous one.", 18, True, GOLD),
], align=PP_ALIGN.CENTER)
c2 = add_round(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(3.4),
               NAVY_CARD, GOLD, 2.5)
card_text(c2, [
    ("THE FEED", 16, True, GOLD),
    ("$2.1 billion", 36, True, WHITE),
    ("in fraud losses started on\nsocial media (FTC 2025)", 18, False, LIGHT),
    ("", 10, False, WHITE),
    ("Facebook beat text and email\nas the starting point for volume.", 18, True, GOLD),
], align=PP_ALIGN.CENTER)
add_textbox(s, "Social media starts more of the volume. The phone still takes the most per person.",
            Inches(0.55), Inches(5.2), Inches(12.23), Inches(0.7),
            font_size=20, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 7. Skill 1 header ───────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
add_textbox(s, "SKILL 1",
            Inches(0.55), Inches(2.15), Inches(12.23), Inches(0.55),
            font_size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Spot a Fake",
            Inches(0.55), Inches(2.7), Inches(12.23), Inches(1.1),
            font_size=56, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "A fake wants you to believe the costume.\nLook at the costume, not the logo.",
            Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.3),
            font_size=24, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 8. Four red flags ───────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "Four red flags to check")
flags = [
    ("THE SENDER", "Would Amazon write from\nemail-secure-alert.com?\nHover or tap. Do not click."),
    ("THE GREETING", "“Dear Customer” from a bank\nthat knows your name\nis a tell. Often."),
    ("THE ASK", "A link to verify. A QR code\nyou did not ask for. A code\nthat just arrived on your phone."),
    ("THE SURPRISE", "You did not start this\nconversation. That matters\nmore than the grammar."),
]
for i, (head, body) in enumerate(flags):
    left = Inches(0.45 + i * 3.2)
    card = add_round(s, left, Inches(1.5), Inches(3.05), Inches(3.7),
                     NAVY_CARD, GOLD, 2)
    add_textbox(s, head,
                left + Inches(0.08), Inches(1.65), Inches(2.9), Inches(0.55),
                font_size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(s, body,
                left + Inches(0.12), Inches(2.25), Inches(2.82), Inches(2.6),
                font_size=16, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "Warning: “No typos” is no longer your safety test. AI fixed the spelling. That test is dead.",
            Inches(0.55), Inches(5.4), Inches(12.23), Inches(0.7),
            font_size=20, bold=True, color=RED, align=PP_ALIGN.CENTER)

# ── 9. Real agencies ────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "What real agencies actually do")
rows = [
    ("Social Security", "does not suspend your number because you missed a text."),
    ("The IRS", "does not demand gift cards."),
    ("The FTC", "will never threaten you, tell you to move money, or send you to a Bitcoin ATM."),
]
for i, (who, what) in enumerate(rows):
    top = Inches(1.5 + i * 1.15)
    card = add_round(s, Inches(0.55), top, Inches(12.23), Inches(1.05),
                     NAVY_CARD, GOLD, 1.5)
    add_textbox(s, who,
                Inches(0.8), top + Inches(0.28), Inches(3.2), Inches(0.5),
                font_size=22, bold=True, color=GOLD)
    add_textbox(s, what,
                Inches(4.1), top + Inches(0.22), Inches(8.4), Inches(0.65),
                font_size=20, color=WHITE)

add_textbox(s, "Bank  ·  Medicare  ·  Amazon  ·  Microsoft  ·  Apple  ·  Sheriff  ·  Grandchild  ·  Tech support that called you",
            Inches(0.55), Inches(5.1), Inches(12.23), Inches(0.7),
            font_size=16, color=LIGHT, align=PP_ALIGN.CENTER)
add_textbox(s, "Treat every one as a costume until YOU started the call on a number you already had.",
            Inches(0.55), Inches(5.7), Inches(12.23), Inches(0.5),
            font_size=18, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

# ── 10. How to check ────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "How to check")
steps = [
    ("1", "HANG UP.", GOLD),
    ("2", "Look up the number on a statement,\nthe back of the card, or the official\nwebsite you type yourself.", WHITE),
    ("3", "CALL THAT NUMBER.", GOLD),
]
for i, (num, text, color) in enumerate(steps):
    left = Inches(0.55 + i * 4.2)
    card = add_round(s, left, Inches(1.55), Inches(3.95), Inches(3.15),
                     NAVY_CARD, GOLD, 2.5)
    add_textbox(s, num,
                left, Inches(1.7), Inches(3.95), Inches(0.7),
                font_size=40, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(s, text,
                left + Inches(0.18), Inches(2.5), Inches(3.6), Inches(1.9),
                font_size=20, bold=(i != 1), color=color, align=PP_ALIGN.CENTER)
add_textbox(s, "Never call the number in the text. That number is the costume’s sleeve.",
            Inches(0.55), Inches(5.0), Inches(12.23), Inches(0.55),
            font_size=22, bold=True, color=RED, align=PP_ALIGN.CENTER)
add_textbox(s, "This is the Federal Trade Commission’s advice. Hang up. Look up. Call back.",
            Inches(0.55), Inches(5.6), Inches(12.23), Inches(0.45),
            font_size=18, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 11. Social media scene ──────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "The newest costume is your feed")
c1 = add_round(s, Inches(0.55), Inches(1.5), Inches(5.9), Inches(4.6),
               NAVY_CARD, GOLD, 2)
card_text(c1, [
    ("FTC 2025", 16, True, GOLD),
    ("$2.1 billion", 36, True, WHITE),
    ("lost — started on social media", 18, False, LIGHT),
    ("", 8, False, WHITE),
    ("Facebook beat text and email\nas the starting point for fraud volume.", 18, False, WHITE),
    ("", 8, False, WHITE),
    ("I am not telling you to close your accounts.\nI am telling you the costume changed.", 16, True, GOLD),
], align=PP_ALIGN.CENTER)
c2 = add_round(s, Inches(6.7), Inches(1.5), Inches(6.1), Inches(4.6),
               NAVY_CARD, GOLD, 2)
card_text(c2, [
    ("Before you click anything in your feed", 16, True, GOLD),
    ("Who started this contact?", 20, True, WHITE),
    ("Does this store or person have a real history?", 20, True, WHITE),
    ("Is there urgency or a prize?", 20, True, WHITE),
    ("", 8, False, WHITE),
    ("Fake ads. Fake profiles. Taken-over DMs.\nFake giveaways.", 16, False, LIGHT),
    ("", 6, False, WHITE),
    ("Type the real site yourself.\nCall the number you already have.\nDo not click the link in the post.", 16, True, GOLD),
], align=PP_ALIGN.LEFT)

# ── 12. Skill 2 header ──────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
add_textbox(s, "SKILL 2",
            Inches(0.55), Inches(2.15), Inches(12.23), Inches(0.55),
            font_size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Name the Pressure",
            Inches(0.55), Inches(2.7), Inches(12.23), Inches(1.1),
            font_size=56, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "Fakes work because of feelings, not because you cannot read.",
            Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.0),
            font_size=24, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 13. Four levers ─────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "The four levers. Say them out loud.")
levers = [
    ("HURRY", "Twenty minutes. A warrant.\nThe refund expires.\nIf you are not allowed to think,\nyou are not in a real process."),
    ("FEAR", "Arrest. Deportation. A virus.\nA stolen identity they will “fix.”\nFear makes smart people obedient."),
    ("SECRECY", "Don’t tell your wife.\nShe would worry.\nIsolation is the scam.\nA real emergency can stand a second adult."),
    ("AUTHORITY", "A badge. A case number.\nCaller ID that says IRS.\nCaller ID is a costume.\nThey can fake it."),
]
for i, (head, body) in enumerate(levers):
    left = Inches(0.45 + i * 3.2)
    card = add_round(s, left, Inches(1.5), Inches(3.05), Inches(4.6),
                     NAVY_CARD, RED, 2.5)
    add_textbox(s, head,
                left + Inches(0.08), Inches(1.7), Inches(2.9), Inches(0.6),
                font_size=22, bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_textbox(s, body,
                left + Inches(0.14), Inches(2.4), Inches(2.78), Inches(3.4),
                font_size=16, color=WHITE, align=PP_ALIGN.CENTER)

# ── 14. Never move money ────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "The fifth lever. Memorize it.")
add_textbox(s, "Never move money to protect it.",
            Inches(0.55), Inches(1.4), Inches(12.23), Inches(0.8),
            font_size=36, bold=True, color=RED, align=PP_ALIGN.CENTER)
add_textbox(s, "Someone says your account is compromised. To keep the money “safe,”\n"
               "you must wire it, buy gold, feed a crypto ATM, or hand cash to a courier.",
            Inches(0.7), Inches(2.25), Inches(11.9), Inches(1.0),
            font_size=20, color=LIGHT, align=PP_ALIGN.CENTER)
add_textbox(s, "That money is not going into a vault with your name on it. It is going to them.",
            Inches(0.7), Inches(3.3), Inches(11.9), Inches(0.55),
            font_size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
stop = [
    "Gift cards",
    "Wire transfers",
    "Crypto / Bitcoin ATM",
    "Courier at the door",
    "Remote access",
]
for i, label in enumerate(stop):
    left = Inches(0.45 + i * 2.56)
    card = add_round(s, left, Inches(4.1), Inches(2.42), Inches(1.5),
                     NAVY_CARD, RED, 2)
    add_textbox(s, label,
                left + Inches(0.06), Inches(4.4), Inches(2.3), Inches(1.0),
                font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "No legitimate tax office, bank, or grandchild-in-trouble story needs a Walmart gift card.",
            Inches(0.55), Inches(5.75), Inches(12.23), Inches(0.4),
            font_size=16, italic=True, color=GOLD, align=PP_ALIGN.CENTER)

# ── 15. Tech support ────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "If you did not call them, they do not get to drive")
c1 = add_round(s, Inches(0.55), Inches(1.5), Inches(12.23), Inches(2.2),
               NAVY_CARD, GOLD, 2)
card_text(c1, [
    ("TECH SUPPORT", 16, True, GOLD),
    ("A pop-up says your computer is infected. A polite person offers to take over the screen.", 22, False, WHITE),
    ("Older adults are five times more likely than younger adults to report losing money on that scam (FTC).", 18, False, LIGHT),
], align=PP_ALIGN.LEFT)
c2 = add_round(s, Inches(0.55), Inches(3.9), Inches(12.23), Inches(2.15),
               NAVY_CARD, GOLD, 2)
card_text(c2, [
    ("AARP still sees more impostors posing as companies and agencies than as family.", 20, True, WHITE),
    ("Do not only guard the grandchild story. Guard the Amazon + bank + sheriff pile-up too.", 20, False, LIGHT),
    ("They will transfer you from one costume to the next in a single call.", 18, True, GOLD),
], align=PP_ALIGN.LEFT)

# ── 16. AI voice / code word ────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "The new costume — a voice you love")
add_textbox(s, "A familiar voice is no longer proof.",
            Inches(0.55), Inches(1.4), Inches(12.23), Inches(0.7),
            font_size=32, bold=True, color=RED, align=PP_ALIGN.CENTER)
add_textbox(s, "AARP: 8 or 9 in 10 adults 50+ worry about cloned voices.\n"
               "AI-flagged scams jumped about 20× from 2023 to 2025.",
            Inches(0.7), Inches(2.15), Inches(11.9), Inches(1.0),
            font_size=20, color=LIGHT, align=PP_ALIGN.CENTER)
c = add_round(s, Inches(1.4), Inches(3.3), Inches(10.5), Inches(2.7),
              NAVY_CARD, GREEN, 2.5)
card_text(c, [
    ("21st-century trick.  20th-century fix.", 16, True, GOLD),
    ("Pick a family code word.", 32, True, GREEN),
    ("Known only to your people. Used when money or panic shows up.", 18, False, WHITE),
    ("If they don’t know it — hang up. Then call the number in your phone.", 18, True, GOLD),
], align=PP_ALIGN.CENTER)

# ── 17. Skill 3 header ──────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
add_textbox(s, "SKILL 3",
            Inches(0.55), Inches(2.15), Inches(12.23), Inches(0.55),
            font_size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Passwords & the Second Lock",
            Inches(0.55), Inches(2.7), Inches(12.23), Inches(1.2),
            font_size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "A password is a key. Most people have one key and twenty doors.",
            Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.0),
            font_size=24, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 18. Password rules ──────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "NIST’s updated password rules")
rules = [
    ("LONGER", "Aim for 15 or more characters.\nA phrase beats a scrambled string."),
    ("KINDER", "maple-porch-radio-Thursday\nis stronger than P@ssw0rd!"),
    ("UNIQUE", "Email cannot match the bank.\nEmail is the skeleton key."),
    ("STABLE", "You do not have to change it\nevery 90 days."),
]
for i, (head, body) in enumerate(rules):
    col = i % 2
    row = i // 2
    left = Inches(0.55 + col * 6.35)
    top = Inches(1.5 + row * 2.35)
    card = add_round(s, left, top, Inches(6.1), Inches(2.15),
                     NAVY_CARD, GOLD, 2)
    add_textbox(s, head,
                left + Inches(0.25), top + Inches(0.2), Inches(5.6), Inches(0.5),
                font_size=20, bold=True, color=GOLD)
    add_textbox(s, body,
                left + Inches(0.25), top + Inches(0.75), Inches(5.6), Inches(1.15),
                font_size=20, color=WHITE)

# ── 19. Two honest options ──────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "If you cannot remember twenty phrases")
c1 = add_round(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(4.5),
               NAVY_CARD, GOLD, 2.5)
card_text(c1, [
    ("OPTION 1", 16, True, GOLD),
    ("A password manager", 28, True, WHITE),
    ("A vault app — if you have someone\nwho can sit with you once to set it up.", 18, False, LIGHT),
    ("Protect that vault with a long phrase\nand a second lock.", 18, False, WHITE),
], align=PP_ALIGN.CENTER)
c2 = add_round(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(4.5),
               NAVY_CARD, GOLD, 2.5)
card_text(c2, [
    ("OPTION 2", 16, True, GOLD),
    ("A small notebook", 28, True, WHITE),
    ("In a drawer or a lockbox at home.\nNot a sticky note on the monitor.\nNot a photo in an unlocked camera roll.", 18, False, LIGHT),
    ("A hidden notebook beats one reused\npassword on the internet.", 18, True, GOLD),
], align=PP_ALIGN.CENTER)

# ── 20. 2FA ─────────────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "Turn on the second lock this week")
add_textbox(s, "Two-factor authentication  =  password  +  something you have",
            Inches(0.55), Inches(1.4), Inches(12.23), Inches(0.55),
            font_size=22, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
c1 = add_round(s, Inches(0.55), Inches(2.1), Inches(6.0), Inches(3.95),
               NAVY_CARD, GOLD, 2)
card_text(c1, [
    ("WHERE", 16, True, GOLD),
    ("Settings  →  Security", 24, True, WHITE),
    ("Look for Two-factor,\n2-step verification, or Multifactor.", 18, False, LIGHT),
    ("The words change. The idea does not.", 18, True, GOLD),
    ("", 8, False, WHITE),
    ("Email first. Then the bank.", 22, True, GREEN),
], align=PP_ALIGN.CENTER)
c2 = add_round(s, Inches(6.8), Inches(2.1), Inches(6.0), Inches(3.95),
               NAVY_CARD, GOLD, 2)
card_text(c2, [
    ("HOW STRONG", 16, True, GOLD),
    ("CISA: any second factor\nis better than none.", 20, True, WHITE),
    ("An authenticator app or a passkey\nis stronger than a text message.", 18, False, LIGHT),
    ("Text messages can be stolen.", 18, False, WHITE),
    ("", 6, False, WHITE),
    ("Perfect is the enemy of a locked door.", 18, True, GOLD),
], align=PP_ALIGN.CENTER)

# ── 21. Never read a code ───────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "Never read a code to anyone")
add_textbox(s, "The code that just arrived is a second key.",
            Inches(0.55), Inches(1.55), Inches(12.23), Inches(0.7),
            font_size=28, color=WHITE, align=PP_ALIGN.CENTER)
add_round(s, Inches(1.3), Inches(2.4), Inches(10.7), Inches(2.5),
          NAVY_CARD, RED, 3)
add_textbox(s, "Anyone who calls or texts and asks you to read it\n"
               "is a scammer — even if they already know your name,\n"
               "even if they say they are the bank.",
            Inches(1.5), Inches(2.7), Inches(10.3), Inches(2.0),
            font_size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "You type the code into the site YOU opened. You never speak it to a stranger.",
            Inches(0.55), Inches(5.2), Inches(12.23), Inches(0.7),
            font_size=24, bold=True, color=GOLD, align=PP_ALIGN.CENTER)

# ── 22. If it already happened ──────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "If it already happened")
add_textbox(s, "This is not a moral failure. You are still a competent person.",
            Inches(0.55), Inches(1.35), Inches(12.23), Inches(0.5),
            font_size=20, italic=True, color=GREEN, align=PP_ALIGN.CENTER)
steps = [
    ("1", "Stop sending money."),
    ("2", "Call the bank — the number on the card. Say wire, gift card, or crypto ATM."),
    ("3", "Change your email password and turn on two-factor."),
    ("4", "Report it: ReportFraud.ftc.gov  and  ic3.gov"),
    ("5", "Tell one trusted person."),
]
for i, (num, text) in enumerate(steps):
    top = Inches(1.95 + i * 0.78)
    add_round(s, Inches(0.7), top, Inches(0.7), Inches(0.62), NAVY_CARD, GOLD, 1.5)
    add_textbox(s, num,
                Inches(0.7), top + Inches(0.05), Inches(0.7), Inches(0.52),
                font_size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(s, text,
                Inches(1.6), top + Inches(0.08), Inches(10.8), Inches(0.55),
                font_size=20, bold=(i in (0, 3)), color=WHITE)
add_textbox(s, "Speed matters. Recovery is not a promise. Silence is a guarantee of no chance.",
            Inches(0.55), Inches(6.0), Inches(12.23), Inches(0.35),
            font_size=16, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

# ── 23. Homework / fridge ───────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
title_bar(s, "Your homework — twenty minutes")
hw = [
    ("1", "Put the hang-up rule on the fridge."),
    ("2", "Send the family code word to the people who matter. Not posted online."),
    ("3", "Turn on email two-factor. Then the bank."),
    ("4", "Set one unique long phrase for email if it still matches other sites."),
    ("5", "If a scam already happened — file both reports today."),
]
for i, (num, text) in enumerate(hw):
    top = Inches(1.4 + i * 0.85)
    add_round(s, Inches(0.7), top, Inches(0.7), Inches(0.65), NAVY_CARD, GOLD, 1.5)
    add_textbox(s, num,
                Inches(0.7), top + Inches(0.05), Inches(0.7), Inches(0.55),
                font_size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(s, text,
                Inches(1.6), top + Inches(0.08), Inches(10.8), Inches(0.55),
                font_size=20, color=WHITE)
add_textbox(s, "If you do nothing else: the fridge line and the code word.",
            Inches(0.55), Inches(5.75), Inches(12.23), Inches(0.4),
            font_size=18, italic=True, color=GOLD, align=PP_ALIGN.CENTER)

# ── 24. Fridge line ─────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s)
add_textbox(s, "THE FRIDGE LINE",
            Inches(0.55), Inches(1.5), Inches(12.23), Inches(0.45),
            font_size=18, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
card = add_round(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(3.6),
                 NAVY_CARD, GOLD, 3)
add_textbox(s, "UNEXPECTED  +  HURRY  +  MONEY OR A CODE",
            Inches(1.1), Inches(2.5), Inches(11.1), Inches(1.0),
            font_size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "=",
            Inches(1.1), Inches(3.4), Inches(11.1), Inches(0.5),
            font_size=28, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "HANG UP AND VERIFY",
            Inches(1.1), Inches(3.95), Inches(11.1), Inches(0.9),
            font_size=36, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Hang up.  Look up.  Call back.",
            Inches(0.55), Inches(5.95), Inches(12.23), Inches(0.45),
            font_size=22, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

# ── 25. Closing ─────────────────────────────────────────────────────────────
s = blank_slide(prs)
chrome(s, footer="The series is free. Nobody here will ask you for a gift card.")
add_textbox(s, "LESSON 3  |  50+TECHBRIDGE",
            Inches(0.55), Inches(1.15), Inches(12.23), Inches(0.4),
            font_size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "Hang up. Look up. Call back.",
            Inches(0.55), Inches(1.65), Inches(12.23), Inches(1.0),
            font_size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "You can spot a costume. Name the pressure. Lock the door.",
            Inches(0.55), Inches(2.7), Inches(12.23), Inches(0.5),
            font_size=20, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)

add_round(s, Inches(1.6), Inches(3.4), Inches(10.1), Inches(2.45),
          NAVY_CARD, GOLD, 2)
add_textbox(s, "To report a scam",
            Inches(1.8), Inches(3.55), Inches(9.7), Inches(0.4),
            font_size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
add_textbox(s, "ReportFraud.ftc.gov     ·     ic3.gov\n833-FRAUD-11     ·     877-908-3360",
            Inches(1.8), Inches(4.0), Inches(9.7), Inches(1.0),
            font_size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(s, "learnmoretechnologies.com/start-free-lesson",
            Inches(1.8), Inches(5.1), Inches(9.7), Inches(0.4),
            font_size=16, color=LIGHT, align=PP_ALIGN.CENTER)

add_textbox(s, "I’m Brian McKinney.  I’ll see you in the room.",
            Inches(0.55), Inches(6.05), Inches(12.23), Inches(0.4),
            font_size=18, color=MUTED, align=PP_ALIGN.CENTER)

# ── Save ────────────────────────────────────────────────────────────────────
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print(f"Saved: {OUT}")
print(f"Slides: {len(prs.slides)}")

root = os.path.normpath(os.path.join(os.path.dirname(OUT), "..", "..", "..", "Dont Get Scammed.pptx"))
# Workspace root is three levels up from Lesson3? 
# Lesson3 is at: 01-LESSONS/20260818 LMTnew lessons script/Lesson3
# Workspace: LMT-CONTENT
root = os.path.normpath(os.path.join(os.path.dirname(OUT), "..", "..", "..", "Dont Get Scammed.pptx"))
try:
    prs.save(root)
    print(f"Also saved: {root}")
except Exception as e:
    print(f"Root copy skipped: {e}")
