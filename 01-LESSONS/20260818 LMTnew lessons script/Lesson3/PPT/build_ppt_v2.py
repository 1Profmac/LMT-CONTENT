"""
Lesson 3 — Don't Get Scammed  (v2)
50+TechBridge  ·  ~15 slides, one idea each, huge text
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os
import uuid

NAVY  = RGBColor(0x0E, 0x1C, 0x2F)
GOLD  = RGBColor(0xC8, 0x94, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xC8, 0xCD, 0xD6)
RED   = RGBColor(0xE8, 0x5D, 0x4C)
GREEN = RGBColor(0x4C, 0xB8, 0x8A)
FONT  = "Calibri"

W = Inches(13.333)
H = Inches(7.5)

OUT  = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "DontGetScammed-Lesson3-v2.pptx")
ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         "..", "..", "..", "Dont Get Scammed v2.pptx"))


# ── Low-level helpers ────────────────────────────────────────────────────────

def new_prs():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    return prs


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    return s


def rect(s, left, top, w, h, fill, no_line=True):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if no_line:
        sh.line.fill.background()
    return sh


def txt(s, text, left, top, w, h,
        size=32, bold=False, color=WHITE, align=PP_ALIGN.LEFT, italic=False):
    box = s.shapes.add_textbox(left, top, w, h)
    tf  = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size  = Pt(size)
    r.font.bold  = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name  = FONT
    return box


# ── High-level slide builders ────────────────────────────────────────────────

LABEL_TEXT = "LESSON 3  ·  50+TechBridge"

def chrome(s):
    """Top label + gold bars top/bottom."""
    txt(s, LABEL_TEXT,
        Inches(0.5), Inches(0.14), Inches(8), Inches(0.36),
        size=14, bold=True, color=GOLD)
    rect(s, Inches(0.5), Inches(0.54), Inches(12.333), Pt(3), GOLD)
    rect(s, Inches(0.5), Inches(6.94), Inches(12.333), Pt(3), GOLD)


def heading(s, text, color=GOLD, size=34):
    txt(s, text,
        Inches(0.5), Inches(0.65), Inches(12.333), Inches(0.85),
        size=size, bold=True, color=color)


def bullets(s, items, top=Inches(1.6), default_size=32,
            left=Inches(0.65), width=Inches(12.1)):
    """
    items = list of str  OR  (text, size, bold, color)
    """
    box = s.shapes.add_textbox(left, top, width, Inches(5.1))
    tf  = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, str):
            t, sz, bd, col = item, default_size, False, WHITE
        else:
            t   = item[0]
            sz  = item[1] if len(item) > 1 else default_size
            bd  = item[2] if len(item) > 2 else False
            col = item[3] if len(item) > 3 else WHITE
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(16)
        r = p.add_run()
        r.text = t
        r.font.size  = Pt(sz)
        r.font.bold  = bd
        r.font.color.rgb = col
        r.font.name  = FONT


def add_sections(prs, sections):
    """
    Add named sections to the presentation XML.
    sections = list of (name, start_0based_slide_index) tuples
    """
    P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
    URI = '{521415D9-36F7-43E2-AB2F-B90AF26B5E84}'

    prs_el = prs._element

    extLst = prs_el.find(qn('p:extLst'))
    if extLst is None:
        extLst = etree.SubElement(prs_el, qn('p:extLst'))

    # Remove any pre-existing section extension
    for ext in extLst.findall(qn('p:ext')):
        if ext.get('uri') == URI:
            extLst.remove(ext)

    ext = etree.SubElement(extLst, qn('p:ext'))
    ext.set('uri', URI)
    sectionLst = etree.SubElement(ext, f'{{{P14}}}sectionLst')

    # Collect slide IDs in order from presentation XML
    sldIdLst_el = prs_el.find(qn('p:sldIdLst'))
    slide_ids = [int(el.get('id'))
                 for el in sldIdLst_el.findall(qn('p:sldId'))]
    n = len(slide_ids)

    for i, (name, start_idx) in enumerate(sections):
        end_idx = sections[i + 1][1] if i + 1 < len(sections) else n
        sec = etree.SubElement(sectionLst, f'{{{P14}}}section')
        sec.set('name', name)
        sec.set('id', '{' + str(uuid.uuid4()).upper() + '}')
        id_lst = etree.SubElement(sec, f'{{{P14}}}sldIdLst')
        for idx in range(start_idx, end_idx):
            sid = etree.SubElement(id_lst, f'{{{P14}}}sldId')
            sid.set('id', str(slide_ids[idx]))


def skill_slide(prs, num, title):
    """Full-bleed skill header — no bullet content."""
    s = blank(prs)
    chrome(s)
    txt(s, f"SKILL {num}",
        Inches(0.5), Inches(2.4), Inches(12.333), Inches(0.6),
        size=22, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    txt(s, title,
        Inches(0.5), Inches(3.05), Inches(12.333), Inches(1.4),
        size=58, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    return s


# ============================================================================
#  SLIDES  (target ≈ 15)
# ============================================================================

prs = new_prs()

# ── 1. Title ──────────────────────────────────────────────────────────────────
s = blank(prs)
chrome(s)
txt(s, "LESSON 3 OF 3",
    Inches(0.5), Inches(1.5), Inches(12.333), Inches(0.5),
    size=20, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
txt(s, "Don't Get Scammed",
    Inches(0.5), Inches(2.1), Inches(12.333), Inches(1.5),
    size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, "Three skills that keep you in charge",
    Inches(0.5), Inches(3.75), Inches(12.333), Inches(0.7),
    size=28, color=GOLD, align=PP_ALIGN.CENTER)
txt(s, "Frightened people send money.  Clear people hang up.",
    Inches(0.5), Inches(4.6), Inches(12.333), Inches(0.6),
    size=22, italic=True, color=LIGHT, align=PP_ALIGN.CENTER)
txt(s, "Brian McKinney  ·  50+TechBridge",
    Inches(0.5), Inches(6.2), Inches(12.333), Inches(0.5),
    size=18, color=LIGHT, align=PP_ALIGN.CENTER)


# ── 2. You are not the punchline + 3-in-4 ───────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "You are not the punchline.")
# Big stat on the right
txt(s, "3 in 4",
    Inches(7.5), Inches(1.4), Inches(5.3), Inches(2.2),
    size=90, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
txt(s, "older adults who reported a scam\nlost NO money. They spotted it.",
    Inches(7.3), Inches(3.65), Inches(5.5), Inches(1.3),
    size=22, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
# Left side
bullets(s, [
    ("FTC: $2.4 billion lost by adults 60+",      28, False, WHITE),
    ("FBI: Nearly $5 billion lost by over-60s",   28, False, WHITE),
    ("",                                            8, False, WHITE),
    ("Those numbers are ugly.",                    26, False, LIGHT),
    ("They are not a verdict",                     26, False, LIGHT),
    ("on your intelligence.",                      26, True,  GOLD),
], top=Inches(1.6), default_size=28)
# divider
rect(s, Inches(7.0), Inches(1.5), Pt(3), Inches(5.1), GOLD)


# ── 3. Three Skills ──────────────────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "Three skills. That is the whole class.")
bullets(s, [
    ("1.  Spot a Fake",                    44, True, GOLD),
    ("2.  Name the Pressure",              44, True, GOLD),
    ("3.  Passwords & the Second Lock",    44, True, GOLD),
], top=Inches(1.65))


# ── 4. SKILL 1 ───────────────────────────────────────────────────────────────
skill_slide(prs, "1", "Spot a Fake")


# ── 5. Four red flags ────────────────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "Four things to check in every message or call")
bullets(s, [
    ("THE SENDER  —  Does the email address really match the company?",  28, False, WHITE),
    ("THE GREETING  —  'Dear Customer' from your own bank? That's a tell.", 28, False, WHITE),
    ("THE ASK  —  A link? A QR code? A code that just arrived on your phone?", 28, False, WHITE),
    ("THE SURPRISE  —  Did YOU start this conversation?",                28, True,  GOLD),
    ("",                                                                  8, False, WHITE),
    ("'No typos' is dead as a safety test. AI fixed the spelling.",      24, True,  RED),
], top=Inches(1.6), default_size=28)


# ── 6. What agencies do + How to check ──────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "Real agencies vs. the costume")
# Left col — constrained to left half only
bullets(s, [
    ("Social Security does NOT suspend your number.", 26, False, WHITE),
    ("The IRS does NOT demand gift cards.",           26, False, WHITE),
    ("The FTC will NEVER send you to a Bitcoin ATM.", 26, False, WHITE),
    ("",                                              8,  False, WHITE),
    ("Treat every caller as a costume until",        24, False, LIGHT),
    ("YOU started the call on a number you had.",    24, True,  GOLD),
], top=Inches(1.6), default_size=26, left=Inches(0.65), width=Inches(5.9))
# Right col — how to check
rect(s, Inches(7.0), Inches(1.5), Pt(3), Inches(5.1), GOLD)
txt(s, "How to check:",
    Inches(7.3), Inches(1.65), Inches(5.5), Inches(0.55),
    size=22, bold=True, color=GOLD)
txt(s, "1.  Hang up.\n\n2.  Look up the number on your card\n     or the real website you type yourself.\n\n3.  Call THAT number.",
    Inches(7.3), Inches(2.25), Inches(5.5), Inches(3.8),
    size=24, bold=False, color=WHITE)
txt(s, "Never call the number in the text.",
    Inches(7.3), Inches(6.1), Inches(5.5), Inches(0.6),
    size=20, bold=True, color=RED)


# ── 7. SKILL 2 ───────────────────────────────────────────────────────────────
skill_slide(prs, "2", "Name the Pressure")


# ── 8. Four levers ───────────────────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "The four levers. Say them out loud.")
bullets(s, [
    ("HURRY      'You have 20 minutes or there is a warrant.'",  32, True, RED),
    ("FEAR        Arrest. Virus. Stolen identity they will fix.", 32, True, RED),
    ("SECRECY    'Don't tell your wife. She would worry.'",      32, True, RED),
    ("AUTHORITY  Caller ID that says IRS. Caller ID is a costume.", 32, True, RED),
    ("",                                                           10, False, WHITE),
    ("If you are not allowed to think, it is not a real process.", 28, True, GOLD),
], top=Inches(1.6))


# ── 9. Never move money + red flags ─────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "Never move money to protect it.", color=RED, size=36)
txt(s, "That money is not going into a vault. It is going to them.",
    Inches(0.5), Inches(1.6), Inches(12.333), Inches(0.65),
    size=26, color=LIGHT, align=PP_ALIGN.LEFT)
# Red flag payment grid
labels = ["Gift cards", "Wire transfers", "Crypto ATM", "Courier at door", "Remote access"]
for i, lab in enumerate(labels):
    col = i % 3
    row = i // 3
    left = Inches(0.5 + col * 4.15)
    top  = Inches(2.5  + row * 1.6)
    rect(s, left, top, Inches(3.9), Inches(1.35), RGBColor(0x22, 0x3A, 0x58))
    # red left bar
    rect(s, left, top, Inches(0.18), Inches(1.35), RED)
    txt(s, lab,
        left + Inches(0.28), top + Inches(0.28), Inches(3.4), Inches(0.9),
        size=26, bold=True, color=WHITE)
txt(s, "No legitimate process needs a Walmart gift card.",
    Inches(0.5), Inches(6.1), Inches(12.333), Inches(0.55),
    size=20, italic=True, color=GOLD, align=PP_ALIGN.CENTER)


# ── 10. AI voice + code word ─────────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "The new costume — a voice you love")
txt(s, "A familiar voice is no longer proof.",
    Inches(0.5), Inches(1.55), Inches(12.333), Inches(0.75),
    size=34, bold=True, color=RED, align=PP_ALIGN.LEFT)
txt(s, "AI can clone a voice from a few seconds of audio.",
    Inches(0.5), Inches(2.35), Inches(12.333), Inches(0.55),
    size=24, color=LIGHT)
# Code word box
rect(s, Inches(0.5), Inches(3.1), Inches(12.333), Inches(2.85),
     RGBColor(0x16, 0x2C, 0x45))
rect(s, Inches(0.5), Inches(3.1), Inches(12.333), Pt(3), GREEN)
rect(s, Inches(0.5), Inches(5.95), Inches(12.333), Pt(3), GREEN)
txt(s, "21st-century trick.  20th-century fix.",
    Inches(0.75), Inches(3.22), Inches(12.0), Inches(0.45),
    size=18, bold=True, color=GOLD)
txt(s, "Pick a family code word.",
    Inches(0.75), Inches(3.7), Inches(12.0), Inches(0.85),
    size=40, bold=True, color=GREEN)
txt(s, "Use it when money or panic shows up.  If they don't know it — hang up.",
    Inches(0.75), Inches(4.6), Inches(12.0), Inches(0.65),
    size=24, color=WHITE)


# ── 11. SKILL 3 ──────────────────────────────────────────────────────────────
skill_slide(prs, "3", "Passwords & the Second Lock")


# ── 12. Password rules + options ─────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "New password rules (NIST) + two honest options")
# Left — rules
txt(s, "The rules:",
    Inches(0.5), Inches(1.62), Inches(6.3), Inches(0.45),
    size=20, bold=True, color=GOLD)
bullets(s, [
    ("15+ characters — longer is stronger",     26, False, WHITE),
    ("Use a phrase, not a scrambled mess",       26, False, WHITE),
    ("  maple-porch-radio-Thursday",             26, True,  GREEN),
    ("Your email password must be unique",       26, True,  GOLD),
    ("  Email is the skeleton key.",             24, False, LIGHT),
], top=Inches(2.1), default_size=26)
# Divider
rect(s, Inches(6.9), Inches(1.5), Pt(3), Inches(5.1), GOLD)
# Right — two options
txt(s, "If you can't remember them all:",
    Inches(7.1), Inches(1.62), Inches(5.7), Inches(0.45),
    size=20, bold=True, color=GOLD)
txt(s, "OPTION 1\nPassword manager app\n(someone sits with you once to set it up)\n\nOPTION 2\nSmall notebook in a drawer or lockbox\n— NOT on the monitor\n— A hidden notebook beats one reused password",
    Inches(7.1), Inches(2.1), Inches(5.7), Inches(4.5),
    size=22, color=WHITE)


# ── 13. 2FA + Never read a code ──────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "Turn on the second lock — and never read the code aloud")
# Left — 2FA, constrained to left half only
bullets(s, [
    ("Two-factor = password + code sent to your phone", 27, False, WHITE),
    ("Settings  →  Security  →  Two-factor / 2-step",  27, True,  GOLD),
    ("Turn it on for EMAIL first, then the bank.",      27, True,  GREEN),
], top=Inches(1.62), default_size=27, left=Inches(0.65), width=Inches(5.9))
# Divider
rect(s, Inches(6.9), Inches(1.5), Pt(3), Inches(5.1), GOLD)
# Right — never read
txt(s, "Never read a code\nto anyone.",
    Inches(7.1), Inches(1.7), Inches(5.7), Inches(1.7),
    size=38, bold=True, color=RED)
txt(s, "The code is a second key.\nAnyone who asks you to read it\nis a scammer —\neven if they know your name.\n\nYou TYPE it. Never speak it.",
    Inches(7.1), Inches(3.5), Inches(5.7), Inches(3.2),
    size=22, color=WHITE)


# ── 14. If it already happened ───────────────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "If it already happened — this is not a moral failure.")
bullets(s, [
    ("1.  Stop sending money.",                                  34, True,  GOLD),
    ("2.  Call the bank — the number ON the card.",             34, False, WHITE),
    ("3.  Change your email password. Turn on two-factor.",     34, False, WHITE),
    ("4.  Report it:  ReportFraud.ftc.gov  and  ic3.gov",      34, True,  GOLD),
    ("5.  Tell one trusted person.",                             34, False, WHITE),
    ("",                                                         10, False, WHITE),
    ("Speed matters. Silence is a guarantee of no chance.",     24, True,  RED),
], top=Inches(1.62))


# ── 15. Fridge line + homework + closing ─────────────────────────────────────
s = blank(prs)
chrome(s)
heading(s, "Your homework — twenty minutes")
# Fridge line box
rect(s, Inches(0.5), Inches(1.6), Inches(12.333), Inches(1.65),
     RGBColor(0x16, 0x2C, 0x45))
rect(s, Inches(0.5), Inches(1.6),  Inches(12.333), Pt(3), GOLD)
rect(s, Inches(0.5), Inches(3.25), Inches(12.333), Pt(3), GOLD)
txt(s, "UNEXPECTED + HURRY + MONEY OR A CODE  =  HANG UP AND VERIFY",
    Inches(0.75), Inches(1.75), Inches(11.85), Inches(1.35),
    size=30, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
# Homework list
bullets(s, [
    ("1.  Put the fridge line above on the fridge.",            26, True,  WHITE),
    ("2.  Send the family code word to the people who matter.", 26, False, WHITE),
    ("3.  Turn on email two-factor. Then the bank.",            26, False, WHITE),
    ("4.  Make your email password unique and long.",           26, False, WHITE),
    ("5.  If a scam happened — file both reports today.",       26, True,  RED),
], top=Inches(3.4), default_size=26)
txt(s, "Hang up.  Look up.  Call back.   —   Brian McKinney  ·  50+TechBridge",
    Inches(0.5), Inches(6.25), Inches(12.333), Inches(0.5),
    size=18, color=LIGHT, align=PP_ALIGN.CENTER)


# ── Sections (match voiceover structure) ─────────────────────────────────────
# Slide index (0-based):
#  0  Title
#  1  You are not the punchline + 3-in-4
#  2  Three Skills
#  3  SKILL 1 header
#  4  Four red flags
#  5  Real agencies + How to check
#  6  SKILL 2 header
#  7  Four levers
#  8  Never move money + red flags
#  9  AI voice + family code word
# 10  SKILL 3 header
# 11  Password rules + options
# 12  2FA + Never read a code
# 13  If it already happened
# 14  Homework + Fridge line + Closing

add_sections(prs, [
    ("Opening",                          0),   # slides 1-3
    ("Skill 1: Spot a Fake",             3),   # slides 4-6
    ("Skill 2: Name the Pressure",       6),   # slides 7-9
    ("The New Costume",                  9),   # slide 10
    ("Skill 3: Passwords & the Lock",   10),   # slides 11-13
    ("If It Already Happened",          13),   # slide 14
    ("Homework & Closing",              14),   # slide 15
])

# ── Save ──────────────────────────────────────────────────────────────────────
prs.save(OUT)
print(f"Saved: {OUT}")
print(f"Slides: {len(prs.slides)}")
try:
    prs.save(ROOT)
    print(f"Also saved: {ROOT}")
except Exception as e:
    print(f"Root copy skipped: {e}")
